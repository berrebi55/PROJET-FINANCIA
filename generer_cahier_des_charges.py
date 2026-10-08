from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY=RGBColor(0x0F,0x47,0x61); GREY=RGBColor(0x59,0x59,0x59)
d=Document()
s=d.sections[0]; s.page_width=Cm(21); s.page_height=Cm(29.7)
s.left_margin=s.right_margin=Cm(2.2); s.top_margin=Cm(2.2); s.bottom_margin=Cm(2)
st=d.styles['Normal']; st.font.name='Calibri'; st.font.size=Pt(10.5)
st.element.rPr.rFonts.set(qn('w:eastAsia'),'Calibri')
st.paragraph_format.space_after=Pt(6); st.paragraph_format.line_spacing=1.1
for n,sz,sb in (('Heading 1',17,18),('Heading 2',13,12),('Heading 3',11,8)):
    h=d.styles[n]; h.font.name='Calibri'; h.font.size=Pt(sz); h.font.bold=True; h.font.color.rgb=NAVY
    h.element.rPr.rFonts.set(qn('w:ascii'),'Calibri'); h.element.rPr.rFonts.set(qn('w:hAnsi'),'Calibri')
    h.paragraph_format.space_before=Pt(sb); h.paragraph_format.space_after=Pt(6); h.paragraph_format.keep_with_next=True

def shade(cell,hexc):
    tcPr=cell._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd')
    sh.set(qn('w:val'),'clear'); sh.set(qn('w:color'),'auto'); sh.set(qn('w:fill'),hexc); tcPr.append(sh)
def P(text='',bold=False,italic=False,size=None,color=None,align=None,after=None):
    p=d.add_paragraph(); r=p.add_run(text); r.bold=bold; r.italic=italic
    if size: r.font.size=Pt(size)
    if color: r.font.color.rgb=color
    if align: p.alignment=align
    if after is not None: p.paragraph_format.space_after=Pt(after)
    return p
def RP(parts):
    p=d.add_paragraph()
    for t,b in parts:
        r=p.add_run(t); r.bold=b
    return p
def B(text,lead=None):
    p=d.add_paragraph(style='List Bullet'); p.paragraph_format.space_after=Pt(3)
    if lead: p.add_run(lead).bold=True
    p.add_run(text); return p
def H(t,l=1): return d.add_heading(t,level=l)
def T(headers,rows,widths,fs=9):
    t=d.add_table(rows=1,cols=len(headers)); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=''; r=c.paragraphs[0].add_run(h); r.bold=True; r.font.size=Pt(fs); r.font.color.rgb=RGBColor(255,255,255); shade(c,'0F4761')
    trPr=t.rows[0]._tr.get_or_add_trPr(); th=OxmlElement('w:tblHeader'); th.set(qn('w:val'),'true'); trPr.append(th)
    for k,row in enumerate(rows):
        cells=t.add_row().cells
        for i,v in enumerate(row):
            cells[i].text=''; r=cells[i].paragraphs[0].add_run(str(v)); r.font.size=Pt(fs)
            if i==0: r.bold=True
            if k%2==1: shade(cells[i],'EAF1F5')
    tblPr=t._tbl.tblPr; lay=OxmlElement('w:tblLayout'); lay.set(qn('w:type'),'fixed'); tblPr.append(lay)
    for i,gc in enumerate(t._tbl.tblGrid.findall(qn('w:gridCol'))): gc.set(qn('w:w'),str(int(widths[i]*567)))
    for row in t.rows:
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for i,w in enumerate(widths):
            row.cells[i].width=Cm(w)
            for p in row.cells[i].paragraphs: p.paragraph_format.space_after=Pt(2); p.paragraph_format.line_spacing=1.0
    d.add_paragraph().paragraph_format.space_after=Pt(2)
    return t
def cap(t): P(t,italic=True,size=9,color=GREY,after=8)

# footer page number
def field(run,instr):
    for typ,txt in (('begin',None),(None,instr),('end',None)):
        if typ:
            e=OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'),typ); run._r.append(e)
        else:
            e=OxmlElement('w:instrText'); e.set(qn('xml:space'),'preserve'); e.text=txt; run._r.append(e)
fp=s.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=fp.add_run('DA Consulting – Cahier des charges – Banque Financia – Confidentiel   |   Page '); r.font.size=Pt(8.5); r.font.color.rgb=GREY
r2=fp.add_run(); r2.font.size=Pt(8.5); r2.font.color.rgb=GREY; field(r2,'PAGE')

# ---------- COVER
for _ in range(5): P('')
P('LIVRABLE 2',bold=True,size=12,color=GREY,align=WD_ALIGN_PARAGRAPH.CENTER)
P('CAHIER DES CHARGES',bold=True,size=30,color=NAVY,align=WD_ALIGN_PARAGRAPH.CENTER,after=4)
P('Modernisation et Sécurisation du Système d’Information',size=16,color=NAVY,align=WD_ALIGN_PARAGRAPH.CENTER,after=30)
P('Client : Banque Financia',bold=True,size=13,align=WD_ALIGN_PARAGRAPH.CENTER,after=2)
P('Équipe consultants : DA Consulting',size=12,align=WD_ALIGN_PARAGRAPH.CENTER,after=2)
P('Auditeurs : Alan Amsellem & David-Haï Berrebi',size=12,align=WD_ALIGN_PARAGRAPH.CENTER,after=2)
P('Date : Octobre 2026   |   Version : 1.0',size=12,align=WD_ALIGN_PARAGRAPH.CENTER,after=40)
P('Document établi à partir de la Charte Projet v1.0 (septembre 2026). Les hypothèses de dimensionnement signalées par la mention « hypothèse » sont à valider avec le client lors du comité de lancement.',italic=True,size=9.5,color=GREY,align=WD_ALIGN_PARAGRAPH.CENTER)
d.add_page_break()

# ---------- TOC
H('Sommaire',1)
for t_ in ['1. Introduction et cadrage','2. Besoins métier','3. Besoins techniques','4. Besoins en cybersécurité','5. Contraintes','6. Gouvernance, livrables et recette','7. Risques projet et mesures','8. Conclusion']:
    P(t_,size=11.5,after=5)
d.add_page_break()

# ---------- 1
H('1. Introduction et cadrage')
H('1.1 Objet du document',2)
P('Ce cahier des charges formalise l’expression de besoin de la Banque Financia pour le projet de modernisation et de sécurisation de son système d’information. Il traduit les objectifs SMART et les critères de réussite de la Charte Projet en exigences vérifiables (métier, techniques, cybersécurité), précise les contraintes à respecter et sert de base au lancement du projet, à la consultation des fournisseurs et à la recette finale.')
H('1.2 Rappel du contexte',2)
P('Financia est un établissement bancaire de taille intermédiaire (banque de détail et gestion de patrimoine) évoluant dans un environnement très réglementé (DORA, RGPD, ANSSI, ISO 27001, contrôle ACPR). Quatre faiblesses ont été relevées lors du cadrage :')
B(' logs dispersés (pare-feu, serveurs, applications bancaires), non corrélés ; détection d’incident lente ou inexistante ;',lead='Absence de supervision centralisée :')
B(' comptes créés manuellement, pas de processus formel de retrait des accès (départs, mobilités), comptes à privilèges non maîtrisés ;',lead='Gestion des identités fragmentée :')
B(' absence généralisée de MFA sur les applications sensibles ;',lead='Authentification faible :')
B(' preuves insuffisantes pour les auditeurs internes et externes.',lead='Traçabilité insuffisante :')
H('1.3 Objectifs du projet (rappel de la charte)',2)
T(['Réf.','Objectif SMART','Échéance'],[
 ['O1','Réduire de 80 % les comptes administrateurs locaux non justifiés','6 mois'],
 ['O2','MFA pour 100 % des utilisateurs des applications bancaires critiques','12 mois'],
 ['O3','Mise en service du SOC (détection de niveau 1 opérationnelle)','9 mois'],
 ['O4','Délai moyen de détection d’incident (MTTD) < 30 minutes','Fin de projet'],
 ['O5','Recertification des habilitations de 100 % des collaborateurs, tous les 6 mois','Dès le mois 6'],
 ['O6','Avis favorable sans réserve majeure au prochain audit ACPR / interne','Après mise en production'],
],[1.6,11.4,3.4])
H('1.4 Périmètre',2)
RP([('Inclus : ',True),('socle d’identités et d’accès (IAM, MFA, PAM, recertification), supervision de sécurité (SIEM, SOC niveau 1), protection des postes et serveurs (EDR), segmentation réseau des zones sensibles, durcissement et sauvegardes des systèmes concernés, conduite du changement et formation.',False)])
RP([('Exclus : ',True),('refonte applicative du cœur bancaire, remplacement du réseau étendu agences, déploiement d’un SOC de niveau 3 / réponse forensique avancée (hors phase 1), certification ISO 27001 elle-même (le projet prépare la conformité).',False)])
H('1.5 Hypothèses de dimensionnement',2)
P('La charte ne chiffre pas le parc. Les ordres de grandeur suivants sont retenus comme hypothèses de travail et seront confirmés lors de la phase d’initialisation :')
T(['Élément','Hypothèse retenue'],[
 ['Collaborateurs','≈ 1 200 (dont ≈ 900 utilisateurs d’applications bancaires critiques)'],
 ['Sites','Siège + datacenter principal, site de secours, réseau d’agences'],
 ['Postes de travail','≈ 1 300 postes fixes et portables'],
 ['Serveurs','≈ 250 serveurs physiques et virtuels'],
 ['Comptes à privilèges','≈ 150 comptes (administrateurs, DBA, comptes de service, prestataires)'],
],[5,11.4])

# ---------- 2
H('2. Besoins métier')
P('Les besoins métier expriment ce que la banque doit pouvoir faire, indépendamment des solutions techniques. Chaque exigence est identifiée (BM-xx) et priorisée selon MoSCoW (M = indispensable, S = important, C = souhaitable).')
H('2.1 Continuité d’activité',2)
P('Les services de paiement, de consultation de compte et de gestion de portefeuille ne peuvent subir d’interruption prolongée sans impact client et réglementaire (DORA impose la maîtrise de la résilience opérationnelle).')
T(['ID','Besoin','Exigence mesurable','Priorité'],[
 ['BM-01','Maintenir les services bancaires critiques en cas d’incident cyber','Plan de continuité (PCA) et plan de reprise (PRA) couvrant les applications critiques, testés au moins 1 fois par an','M'],
 ['BM-02','Reprendre rapidement après sinistre','RTO ≤ 4 h et RPO ≤ 15 min pour les applications critiques (valeurs cibles à valider par les métiers)','M'],
 ['BM-03','Contenir un incident sans arrêter la banque','Procédure d’isolement d’un poste/segment en < 15 min sans coupure des services non concernés','M'],
 ['BM-04','Gérer les départs et mobilités sans faille','Retrait des accès sous 24 h après un départ ; revue des droits à chaque mobilité','M'],
 ['BM-05','Produire les preuves pour les audits','Journal d’audit horodaté, exportable, conservé ≥ 12 mois en ligne (rétention légale à confirmer)','M'],
],[1.5,4.6,8.2,2.1])
H('2.2 Mobilité',2)
P('Les conseillers, gestionnaires de patrimoine et équipes support travaillent depuis les agences, le siège et à distance. La sécurité ne doit pas dégrader l’expérience au point d’inciter à contourner les contrôles.')
T(['ID','Besoin','Exigence mesurable','Priorité'],[
 ['BM-06','Accès distant sécurisé aux applications métier','Accès via VPN/ZTNA avec MFA obligatoire ; aucun accès direct non authentifié','M'],
 ['BM-07','Authentification simple pour l’utilisateur','Connexion unique (SSO) sur les applications compatibles ; MFA sans friction (notification/biométrie) ; satisfaction > 80 %','S'],
 ['BM-08','Travail depuis un poste ou mobile maîtrisé','Accès conditionnel selon l’état du poste (chiffré, patché, EDR actif) ; mobiles gérés par MDM','S'],
 ['BM-09','Accès des prestataires encadré','Comptes nominatifs temporaires, accès tracés et expirant automatiquement','M'],
],[1.5,4.6,8.2,2.1])
H('2.3 Disponibilité',2)
T(['ID','Besoin','Exigence mesurable','Priorité'],[
 ['BM-10','Disponibilité des applications critiques','≥ 99,9 % en heures ouvrées ; fenêtres de maintenance planifiées hors pics (fin de mois, clôtures)','M'],
 ['BM-11','Disponibilité de l’authentification','Service IAM/MFA en haute disponibilité ; mode dégradé documenté (codes de secours) ; pas de point de défaillance unique','M'],
 ['BM-12','Disponibilité de la supervision','SOC : couverture 24/7 cible ou plage horaire définie contractuellement (voir §4.3)','M'],
 ['BM-13','Déploiement sans interruption de service','Migrations par vagues, avec plan de retour arrière pour chaque vague','S'],
],[1.5,4.6,8.2,2.1])

# ---------- 3
H('3. Besoins techniques')
H('3.1 Réseau',2)
T(['ID','Besoin','Exigence'],[
 ['BT-01','Cartographie et inventaire','Inventaire des flux et des équipements réseau (matrice de flux à jour) avant toute segmentation'],
 ['BT-02','Segmentation en zones','Séparation au minimum : postes utilisateurs, serveurs applicatifs bancaires, administration, DMZ, partenaires/prestataires, invités'],
 ['BT-03','Filtrage inter-zones','Pare-feu nouvelle génération avec règles en « refus par défaut » ; journalisation de tous les refus et des flux sensibles'],
 ['BT-04','Accès distant','Remplacement des accès VPN simples par un accès avec MFA et contrôle de conformité du poste'],
 ['BT-05','Résilience des liens','Redondance des liens et équipements critiques entre datacenter principal et site de secours'],
],[1.5,4.2,10.7])
H('3.2 Serveurs et postes',2)
T(['ID','Besoin','Exigence'],[
 ['BT-06','Annuaire d’identités','Annuaire (Active Directory / équivalent) assaini : comptes inactifs désactivés, groupes revus, délégation d’administration par niveaux (tiering)'],
 ['BT-07','Supprimer les droits administrateur locaux','Retrait des droits admin locaux injustifiés (-80 %), élévation temporaire à la demande via outil dédié'],
 ['BT-08','Durcissement','Référentiel de configuration (CIS/ANSSI) appliqué aux serveurs et postes ; gestion des correctifs avec délais cibles (critiques ≤ 14 jours)'],
 ['BT-09','Inventaire des actifs','Inventaire unique (CMDB) des serveurs, postes et applications, avec propriétaire et criticité'],
 ['BT-10','Capacité d’accueil','Dimensionnement des serveurs d’infrastructure (SIEM, PAM, IAM) avec marge de croissance de 30 %'],
],[1.5,4.2,10.7])
H('3.3 Sauvegardes',2)
T(['ID','Besoin','Exigence'],[
 ['BT-11','Stratégie de sauvegarde','Règle 3-2-1 : 3 copies, 2 supports différents, 1 copie hors site'],
 ['BT-12','Protection contre le rançongiciel','Au moins une copie immuable ou déconnectée (« air-gap ») ; comptes de sauvegarde distincts, protégés par MFA'],
 ['BT-13','Tests de restauration','Test de restauration complet des applications critiques au moins 1 fois par an, trimestriel par échantillon ; résultats consignés'],
 ['BT-14','Sauvegarde de la configuration de sécurité','Sauvegarde des configurations IAM, PAM, SIEM, pare-feu et règles de détection'],
],[1.5,4.2,10.7])
H('3.4 Cloud',2)
P('Le projet ne prévoit pas de migration massive vers le cloud, mais certaines briques (MFA, SIEM, EDR) sont proposées en mode SaaS. Les exigences suivantes encadrent ce recours, en cohérence avec DORA (risque lié aux prestataires TIC tiers) et le RGPD.')
T(['ID','Besoin','Exigence'],[
 ['BT-15','Localisation des données','Hébergement des données de sécurité dans l’Union européenne ; prestataire non soumis à une législation extraterritoriale contraignante (à vérifier)'],
 ['BT-16','Qualification du prestataire','Préférence pour les offres qualifiées SecNumCloud ou équivalentes ; certification ISO 27001 exigée'],
 ['BT-17','Réversibilité','Clause contractuelle de réversibilité et d’export des données et journaux, sans surcoût caché'],
 ['BT-18','Sécurité de l’intégration','Connecteurs entre cloud et SI interne chiffrés, authentifiés, limités au strict nécessaire'],
 ['BT-19','Registre DORA','Inscription des prestataires cloud au registre d’information des prestataires TIC et clauses contractuelles DORA'],
],[1.5,4.2,10.7])

# ---------- 4
H('4. Besoins en cybersécurité')
H('4.1 MFA et gouvernance des identités (IAM / PAM)',2)
P('C’est le cœur des objectifs O1, O2 et O5. L’authentification forte réduit fortement le risque de vol d’identifiants ; la gouvernance des identités garantit le moindre privilège dans la durée.')
T(['ID','Besoin','Exigence'],[
 ['BC-01','MFA généralisé','MFA obligatoire pour 100 % des utilisateurs des applications bancaires critiques, tous les accès distants et tous les comptes à privilèges'],
 ['BC-02','Facteurs acceptés','Application d’authentification, clé FIDO2 pour comptes sensibles ; SMS toléré uniquement en secours transitoire'],
 ['BC-03','Cycle de vie des identités','Processus automatisé arrivée / mobilité / départ, alimenté par le référentiel RH (source de vérité)'],
 ['BC-04','Coffre-fort à privilèges (PAM)','100 % des comptes à privilèges dans un coffre ; mots de passe aléatoires et rotatifs ; sessions d’administration enregistrées'],
 ['BC-05','Recertification','Campagne de revue des habilitations tous les 6 mois, avec validation par les managers/propriétaires d’application et traçabilité'],
 ['BC-06','Séparation des tâches','Matrice des incompatibilités de droits (ex. initier et valider un paiement) contrôlée'],
],[1.5,4.2,10.7])
H('4.2 EDR et SIEM',2)
T(['ID','Besoin','Exigence'],[
 ['BC-07','EDR sur postes et serveurs','Agent EDR déployé sur ≥ 98 % des postes et serveurs ; blocage automatique des comportements malveillants ; isolement à distance d’un poste'],
 ['BC-08','Collecte centralisée des logs (SIEM)','Collecte des journaux : pare-feu, annuaire, serveurs, EDR, IAM/PAM, VPN, messagerie et applications bancaires critiques'],
 ['BC-09','Corrélation et détection','Cas d’usage de détection (≥ 30 à la mise en service), alignés sur MITRE ATT&CK : brute force, élévation de privilèges, mouvement latéral, exfiltration'],
 ['BC-10','Horodatage et intégrité','Sources synchronisées (NTP) ; journaux protégés contre la modification ; rétention conforme (§5.4)'],
 ['BC-11','Tableaux de bord','Tableaux de bord de suivi des KPI (MTTD, MTTR, couverture MFA/EDR/PAM) pour la direction et le RSSI'],
],[1.5,4.2,10.7])
H('4.3 SOC',2)
T(['ID','Besoin','Exigence'],[
 ['BC-12','Modèle d’exploitation','SOC mutualisé ou externalisé (MSSP) en phase 1, avec option d’internalisation progressive ; choix à arbitrer (voir §5.1)'],
 ['BC-13','Couverture horaire','Cible 24/7 pour la qualification des alertes critiques ; à défaut, plage définie contractuellement et astreinte'],
 ['BC-14','Niveaux de service','MTTD < 30 min ; MTTR < 4 h pour incidents critiques ; accusé de réception d’alerte critique < 15 min'],
 ['BC-15','Processus de réponse à incident','Playbooks documentés (rançongiciel, compte compromis, fuite de données) ; matrice d’escalade ; exercice de crise annuel'],
 ['BC-16','Obligations de notification','Procédure de notification des incidents majeurs (ACPR/autorités compétentes sous DORA, CNIL en cas de violation de données)'],
],[1.5,4.2,10.7])
H('4.4 Segmentation réseau',2)
T(['ID','Besoin','Exigence'],[
 ['BC-17','Isolement des actifs critiques','Zone dédiée aux serveurs bancaires et aux bases de données ; accès uniquement via bastion PAM'],
 ['BC-18','Cloisonnement de l’administration','Réseau d’administration séparé du réseau utilisateur ; postes d’administration dédiés'],
 ['BC-19','Micro-segmentation progressive','Démarrage par les applications les plus sensibles ; mode « observation » avant blocage pour éviter les ruptures de service'],
 ['BC-20','Contrôle des accès tiers','Prestataires isolés en zone dédiée, accès temporaires et supervisés'],
],[1.5,4.2,10.7])

# ---------- 5
H('5. Contraintes')
H('5.1 Budget',2)
P('Le budget n’a pas été communiqué dans la charte. L’enveloppe ci-dessous est une estimation d’ordre de grandeur (hypothèse) destinée à cadrer les arbitrages ; elle sera affinée après consultation des fournisseurs.',after=4)
T(['Poste','Estimation (hypothèse)','Nature'],[
 ['IAM / MFA / PAM (licences et intégration)','250 – 350 k€','Investissement + abonnement'],
 ['SIEM et plate-forme de collecte','150 – 250 k€','Investissement + abonnement'],
 ['EDR (≈ 1 550 agents)','80 – 120 k€ / an','Abonnement'],
 ['SOC externalisé (niveau 1, plage cible 24/7)','200 – 350 k€ / an','Fonctionnement'],
 ['Segmentation réseau et pare-feu','150 – 250 k€','Investissement'],
 ['Prestations de projet (conseil, intégration, recette)','250 – 350 k€','Prestation'],
 ['Formation et conduite du changement','30 – 50 k€','Prestation'],
 ['Provision pour aléas (10 – 15 %)','≈ 150 k€','Réserve'],
 ['Total indicatif première année','≈ 1,3 – 1,9 M€',''],
],[7.6,4.6,4.2])
P('Contrainte de pilotage : tout dépassement supérieur à 10 % de l’enveloppe validée doit être arbitré en comité de pilotage. Les arbitrages SOC externalisé / internalisé et solution SaaS / hébergée sont les principaux leviers de coût.')
H('5.2 Planning',2)
P('Le projet dure 12 mois et doit respecter les jalons de la charte. Le déploiement évite les périodes de forte activité (clôtures trimestrielles, fin d’année).')
T(['Phase','Période','Contenu principal','Jalon'],[
 ['0. Initialisation','M1','Comité de lancement, validation des hypothèses, inventaires, cartographie des flux','Validation du cadrage'],
 ['1. Fondations IAM','M1 – M3','Assainissement de l’annuaire, choix de la solution IAM/PAM, pilote MFA sur les populations à risque (administrateurs)','Pilote MFA validé'],
 ['2. Réduction des privilèges','M2 – M6','Retrait des droits admin locaux, coffre PAM, premiers comptes à privilèges couverts','O1 : -80 % au M6'],
 ['3. Supervision','M3 – M9','Choix SIEM et SOC, collecte des logs, cas d’usage, EDR, mise en service SOC','O3 : SOC niveau 1 au M9'],
 ['4. Généralisation MFA et segmentation','M5 – M11','MFA par vagues, segmentation des zones sensibles, première campagne de recertification','O5 : recertification dès M6'],
 ['5. Recette et audit','M10 – M12','Tests, exercices de crise, mesure des KPI, préparation à l’audit ACPR','O2 : MFA 100 % au M12'],
],[3.4,2.2,7.4,3.4],fs=8.5)
H('5.3 Effectifs',2)
T(['Rôle','Équipe','Charge indicative'],[
 ['Pilotage projet (DA Consulting)','Chef de projet','100 % sur 12 mois'],
 ['Architecture et sécurité (DA Consulting)','Consultants IAM / SOC / réseau (2 à 3)','100 % selon phases'],
 ['Sponsor et gouvernance (Financia)','Direction Générale, RSSI','Comités mensuels ; RSSI ≈ 40 %'],
 ['Exploitation et intégration (Financia)','DSI : équipes systèmes, réseau, support','≈ 3 à 4 ETP à temps partiel'],
 ['Conformité et risques (Financia)','Conformité, DPO, Risques','Revues ponctuelles'],
 ['Métiers','Référents de chaque direction','Tests, recertification, relais du changement'],
],[5.4,6.2,4.8])
P('Risque principal : la disponibilité de la DSI, déjà sollicitée par l’exploitation courante. Un référent DSI nommé et un engagement de temps formalisé en comité de pilotage sont des conditions de démarrage.')
H('5.4 Réglementations et normes',2)
T(['Cadre','Exigences structurantes pour le projet'],[
 ['DORA','Gestion des risques TIC, notification des incidents majeurs, tests de résilience, maîtrise des prestataires TIC tiers (cloud, SOC)'],
 ['RGPD','Minimisation et sécurité des données, notification des violations sous 72 h, analyse d’impact pour la supervision des utilisateurs, information du personnel'],
 ['Recommandations ANSSI','Hygiène informatique, administration sécurisée, journalisation, cloisonnement'],
 ['ISO 27001','Gestion des accès, journalisation, gestion des incidents, amélioration continue'],
 ['Contrôle ACPR / audit interne','Preuves de mise en œuvre : politiques, journaux, comptes rendus de recertification, résultats de tests'],
 ['Droit du travail','Information-consultation des représentants du personnel avant la mise en place d’outils de supervision (SIEM, EDR)'],
],[4.4,12])

# ---------- 6
H('6. Gouvernance, livrables et recette')
H('6.1 Gouvernance',2)
B(' mensuel, Direction Générale, RSSI, DSI, Conformité/Risques, DA Consulting : arbitrages, budget, jalons.',lead='Comité de pilotage :')
B(' hebdomadaire, RSSI, DSI, chef de projet : avancement, risques, décisions techniques.',lead='Point de suivi opérationnel :')
B(' revue des KPI, des écarts et des plans d’action ; tableau de bord partagé.',lead='Reporting :')
H('6.2 Livrables attendus',2)
T(['Livrable','Échéance'],[
 ['Cahier des charges (ce document) et plan de projet détaillé','M1'],
 ['Dossier d’architecture cible (IAM, PAM, SIEM, segmentation) et matrice de flux','M2'],
 ['Politiques : gestion des accès, mots de passe/MFA, journalisation, gestion des incidents','M3'],
 ['Plateforme SIEM/SOC en service, playbooks, cas d’usage de détection','M9'],
 ['Rapports de recertification, de tests PRA et d’exercice de crise','M6 à M12'],
 ['Dossier de preuves d’audit et rapport final de projet avec mesure des KPI','M12'],
],[12,4.4])
H('6.3 Critères de recette',2)
P('La recette s’appuie sur les KPI de la charte, mesurés selon des modalités définies à l’avance :')
T(['KPI (charte)','Cible','Exigence liée','Méthode de vérification'],[
 ['Couverture MFA – applications critiques','100 %','BC-01','Extraction IAM : utilisateurs MFA actif / total'],
 ['Réduction des comptes admin locaux','-80 %','BT-07','Comparaison audit initial / audit final'],
 ['MTTD','< 30 min','BC-14','Alertes de test (exercices) + statistiques SIEM'],
 ['MTTR','< 4 h','BC-14, BC-15','Exercice de crise et incidents réels'],
 ['Couverture PAM','100 %','BC-04','Inventaire des comptes à privilèges / comptes dans le coffre'],
 ['Revues d’habilitations','Tous les 6 mois','BC-05','Comptes rendus de campagnes signés'],
 ['Incidents critiques non détectés','0','BC-08, BC-09','Tests d’intrusion / simulation d’attaque (purple team)'],
 ['Disponibilité du SOC','24/7 ou plage contractuelle','BC-13','Rapports d’activité du prestataire'],
 ['Réserves majeures ACPR','0','Ensemble','Rapport d’audit'],
 ['Satisfaction utilisateurs MFA','> 80 %','BM-07','Enquête post-déploiement'],
],[5,2.4,2.8,6.2],fs=8.5)

# ---------- 7
H('7. Risques projet et mesures')
T(['Risque','Probabilité','Impact','Mesure de maîtrise'],[
 ['Résistance des utilisateurs au MFA','Moyenne','Élevé','Pilote, communication, formation, support renforcé, MFA sans friction'],
 ['Blocage de flux métier lors de la segmentation','Moyenne','Élevé','Mode observation avant blocage, tests en pré-production, retour arrière'],
 ['Indisponibilité des équipes DSI','Élevée','Moyen','Référent nommé, planification des charges, appui DA Consulting'],
 ['Volume d’alertes SOC trop élevé (faux positifs)','Moyenne','Moyen','Réglage progressif des règles, montée en charge par sources'],
 ['Dérive du budget ou des délais','Moyenne','Élevé','Lotissement, jalons de décision, provision pour aléas'],
 ['Compatibilité d’applications bancaires anciennes avec MFA/SSO','Élevée','Moyen','Inventaire préalable, passerelles, protection par PAM/bastion'],
 ['Dépendance à un prestataire SOC/cloud','Faible','Élevé','Clauses DORA, réversibilité, plan de sortie'],
],[5.2,2.2,1.8,7.2],fs=8.5)

H('8. Conclusion')
P('Ce cahier des charges définit 13 exigences métier, 19 exigences techniques et 20 exigences de cybersécurité, rattachées aux six objectifs SMART et aux dix KPI de la charte. Il constitue la référence contractuelle de la consultation, du pilotage et de la recette. Les hypothèses de parc, de budget et de RTO/RPO devront être confirmées dès le comité de lancement afin de sécuriser le planning de 12 mois.')

d.save('Cahier_des_charges_Financia.docx')
